"""TẠO TẬP GÁN NHÃN THỦ CÔNG để kiểm chứng nhãn LLM.

300 bài, chọn NGẪU NHIÊN PHÂN TẦNG theo nhóm tin × dấu nhãn LLM (21 tầng), phân bổ theo căn bậc hai cỡ tầng để lớp
hiếm (tin âm, pháp lý) có đủ mẫu; GHI LẠI xác suất chọn π của từng bài (để hiệu chỉnh sai lệch khi ước lượng cho toàn bộ
kho — Egami et al., 2023). 100 bài được gán bởi 2 người độc lập (đo Cohen's κ). Nhãn LLM bị ẩn khỏi file gán.

Đầu ra: data/manual/team_sheet.xlsx (giao cho nhóm) | data/manual/team_key.parquet (đáp án LLM + π — KHÔNG giao)
Dùng:  python -m src.nlp.manual_sample
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.config import CLEAN, NHOM_TIN, ROOT

N_TOTAL, N_OVERLAP, SEED = 300, 100, 2026
GOLD = ROOT / "data" / "manual"

RELEVANCE = {
    "SU_KIEN_DN": "Tin sự kiện của MỘT/VÀI doanh nghiệp cụ thể (kết quả kinh doanh, phát hành, M&A, pháp lý, nhân sự, giao dịch nội bộ, dự án của chính DN…)",
    "BINH_LUAN_TT": "Bình luận / tổng hợp thị trường chứng khoán (diễn biến VN-Index, nhóm ngành, khối ngoại, top mã tăng giảm…)",
    "CHINH_SACH_NGANH": "Tin chính sách, pháp luật, vĩ mô, thị trường BĐS nói chung — không xoay quanh một DN cụ thể",
    "QUANG_BA": "Bài quảng bá / PR dự án, giới thiệu sản phẩm BĐS, nội dung tài trợ",
    "KHAC": "Không thuộc các loại trên hoặc không liên quan BĐS / chứng khoán",
}
# Định nghĩa loại sự kiện và thang điểm
# đồng nhất với prompt gán nhãn hiện tại trong
# src/label/classify_all.py.
EVENT_TYPES = {
    "LEGAL": "pháp lý, phê duyệt, tranh chấp, thu hồi của dự án cụ thể.",
    "PROJECT": "khởi công, mở bán, bàn giao, tiến độ dự án.",
    "EARNINGS": "kết quả kinh doanh, báo cáo tài chính.",
    "DEBT_BOND": "nợ, thanh toán trái phiếu, đáo hạn, vỡ nợ.",
    "BANKING_CREDIT": "tín dụng ngân hàng, cho vay.",
    "INTEREST_RATE": "thay đổi lãi suất.",
    "POLICY": "luật, quy định, chính sách, quy hoạch nhà nước.",
    "M_AND_A": "mua bán, sáp nhập, chuyển nhượng.",
    "CAPITAL_RAISE": "huy động vốn, phát hành chứng khoán.",
    "INSIDER_TRADING": "giao dịch cổ đông lớn, nội bộ.",
    "MANAGEMENT": "thay đổi lãnh đạo.",
    "MARKET": "cung cầu, giá, diễn biến thị trường chung.",
    "ANALYST": "khuyến nghị, phân tích của tổ chức phân tích.",
    "OTHER": "sự kiện khác.",
    "KHONG_LIEN_QUAN": "bài không liên quan đến bất động sản Việt Nam hoặc doanh nghiệp mục tiêu (khi đó SẮC THÁI = 0).",
}
SENTIMENT = {
    "-2": "tiêu cực mạnh, có căn cứ rõ về ảnh hưởng lớn.",
    "-1": "tiêu cực, có cơ chế tác động cụ thể nhưng chưa đủ căn cứ mức -2.",
    "0": "trung lập, tác động trái chiều cân bằng, hoặc thiếu căn cứ xác định hướng.",
    "1": "tích cực, có cơ chế tác động cụ thể nhưng chưa đủ căn cứ mức 2.",
    "2": "tích cực mạnh, có căn cứ rõ về ảnh hưởng lớn.",
}
# Các quy tắc chấm điểm trong prompt gốc (mục 2, 3, 7) — rút gọn, giữ nguyên ý.
PROMPT_RULES = [
    "Chỉ dùng thông tin trong bài tại thời điểm đăng; không dùng diễn biến giá hay sự kiện xảy ra sau ngày đăng.",
    "Đánh giá hướng tác động KINH TẾ có thể có đối với doanh nghiệp mục tiêu — không phải dự báo giá cổ phiếu.",
    "Không chọn tích cực chỉ vì văn phong lạc quan hay bài mang tính quảng bá; không chọn tiêu cực chỉ vì văn phong bi quan.",
    "Không mặc định dự án lớn, doanh thu lớn hay hợp đồng lớn là tích cực mạnh khi chưa rõ lợi ích, chi phí, quy mô tương đối.",
    "Phân biệt kết quả đã xảy ra với dự kiến, đề xuất, tin đồn. Nhận định của chuyên gia không tự động là sự kiện tích cực.",
    "Nhiều thông tin trái chiều: cân nhắc sự kiện chính và mức trọng yếu. Thiếu dữ kiện xác định hướng: chọn 0.",
    "Từ thiện, giải thưởng, PR thông thường có thể là 0. Mã chỉ được nhắc tên, chưa có tác động xác định: chọn 0.",
    "Không suy ra doanh nghiệp hưởng lợi chỉ vì thuộc ngành BĐS; không gán lợi ích của đối thủ / công ty mẹ / con cho mã.",
    "Không tự diễn giải 'gặp khó khăn' thành 'thua lỗ nặng', 'cơ cấu lại nợ' thành 'nguy cơ vỡ nợ', "
    "'chậm tiến độ' thành 'đình trệ hoàn toàn'. Không chấm −2 chỉ vì có từ khóa tiêu cực.",
]


def llm_article_labels():
    """Nhãn LLM cấp bài: ưu tiên nhãn cấp ngành (keyword/law); bài chỉ có ở company/ -> event_type phổ biến nhất, điểm TB."""
    lab = pd.read_parquet(CLEAN / "article_label.parquet")
    ind = lab[lab["typ"].isin(["keyword", "law"])].groupby("article_id").agg(
        llm_event_type=("event_type", lambda s: s.mode().iat[0]), llm_score=("impact_score", "mean"))
    atl = pd.read_parquet(CLEAN / "article_ticker_label.parquet")
    com = atl.groupby("article_id").agg(llm_event_type=("event_type", lambda s: s.mode().iat[0]),
                                        llm_score=("impact_score", "mean"), tickers=("ticker", lambda s: ", ".join(sorted(set(s)))))
    out = ind.combine_first(com[["llm_event_type", "llm_score"]])
    out["tickers"] = com["tickers"].reindex(out.index).fillna("")
    out["llm_source"] = np.where(out.index.isin(ind.index), "nganh", "ma")
    return out.reset_index()


def main():
    art = pd.read_parquet(CLEAN / "articles_clean.parquet",
                          columns=["article_id", "src", "published_at", "title", "description", "content", "short_content"])
    d = art[~art["short_content"]].merge(llm_article_labels(), on="article_id")
    d["nhom"] = d["llm_event_type"].map(NHOM_TIN).fillna("khac")
    d["dau"] = np.sign(d["llm_score"].round(6)).astype(int)
    d["stratum"] = d["nhom"] + "|" + d["dau"].astype(str)
    N_h = d.groupby("stratum").size()
    alloc = np.sqrt(N_h); alloc = (alloc / alloc.sum() * N_total_adj(N_h)).round().astype(int).clip(lower=2)
    alloc = np.minimum(alloc, N_h)
    rng = np.random.default_rng(SEED)
    parts = []
    for s, n in alloc.items():
        g = d[d["stratum"] == s]
        idx = rng.choice(len(g), size=int(n), replace=False)
        parts.append(g.iloc[idx].assign(pi=n / len(g)))
    smp = pd.concat(parts).sample(frac=1, random_state=SEED).head(N_TOTAL).reset_index(drop=True)
    smp["gold_id"] = [f"G{i + 1:03d}" for i in range(len(smp))]
    smp["overlap"] = False
    smp.loc[rng.choice(len(smp), size=N_OVERLAP, replace=False), "overlap"] = True
    GOLD.mkdir(parents=True, exist_ok=True)
    key = smp[["gold_id", "article_id", "stratum", "pi", "overlap", "llm_event_type", "llm_score", "llm_source", "src"]]
    key.to_parquet(GOLD / "team_key.parquet", index=False)
    write_xlsx(smp)
    print(f"Tập vàng: {len(smp)} bài ({smp['overlap'].sum()} bài gán 2 người) từ {len(d):,} bài | {len(alloc)} tầng")
    print("Phân bố theo nhóm tin:", smp["nhom"].value_counts().to_dict())
    print("Phân bố theo dấu nhãn LLM:", smp["dau"].value_counts().to_dict())


def N_total_adj(N_h):
    return N_TOTAL + 10          # dư một ít vì làm tròn + chặn dưới 2 bài/tầng; cắt lại còn đúng N_TOTAL


def write_xlsx(smp):
    wb = Workbook()
    hd = wb.active; hd.title = "Huong_dan"
    lines = [("HƯỚNG DẪN GÁN NHÃN TẬP VÀNG", True), ("", False),
             ("• Đọc tiêu đề, mô tả và trích đoạn nội dung; gán 3 nhãn ở 3 cột màu vàng (chọn từ danh sách thả xuống).", False),
             ("• Gán ĐỘC LẬP: không trao đổi với người gán kia, không tra nhãn của LLM.", False),
             ("• Loại sự kiện và sắc thái dùng ĐÚNG định nghĩa trong prompt đã giao cho LLM (qwen2.5:7b) — xem mục 2, 3, 4.", False),
             ("• Sắc thái = hướng tác động kinh tế đối với các mã được nhắc (nếu không có mã thì với ngành BĐS).", False),
             ("• Người A gán sheet Nguoi_A (300 bài). Người B gán sheet Nguoi_B (100 bài — trùng một phần với người A).", False),
             ("• Thời gian ước tính: khoảng 1 phút/bài.", False), ("", False),
             ("1. MỨC LIÊN QUAN (nhãn MỚI, LLM không gán — dùng để kiểm tra bộ lọc tin)", True)]
    lines += [(f"   {k}: {v}", False) for k, v in RELEVANCE.items()]
    lines += [("", False), ("2. LOẠI SỰ KIỆN (chọn đúng một, cho sự kiện CHÍNH của bài)", True)]
    lines += [(f"   {k}: {v}", False) for k, v in EVENT_TYPES.items()]
    lines += [("", False), ("3. SẮC THÁI / ĐIỂM TÁC ĐỘNG (−2 … +2)", True)] + [(f"   {k}: {v}", False) for k, v in SENTIMENT.items()]
    lines += [("", False), ("4. QUY TẮC CHẤM (từ prompt gốc)", True)] + [(f"   • {r}", False) for r in PROMPT_RULES]
    for i, (txt, bold) in enumerate(lines, 1):
        hd.cell(row=i, column=1, value=txt).font = Font(bold=bold, size=12 if bold else 11)
    hd.column_dimensions["A"].width = 150
    cols = ["gold_id", "Nguồn", "Ngày đăng", "Mã được nhắc", "Tiêu đề", "Mô tả", "Trích nội dung",
            "MỨC LIÊN QUAN", "LOẠI SỰ KIỆN", "SẮC THÁI (−2..2)", "Ghi chú"]
    widths = [8, 10, 12, 14, 45, 60, 90, 18, 18, 14, 25]
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for sheet, rows in (("Nguoi_A", smp), ("Nguoi_B", smp[smp["overlap"]].sample(frac=1, random_state=7))):
        ws = wb.create_sheet(sheet)
        ws.append(cols)
        for c, w in zip(ws[1], widths):
            c.font = Font(bold=True); ws.column_dimensions[c.column_letter].width = w
        for r in rows.itertuples():
            ws.append([r.gold_id, r.src, str(r.published_at)[:16], r.tickers, r.title, r.description,
                       (r.content or "")[:800], None, None, None, None])
        n = len(rows) + 1
        for col in "EFG":
            for cell in ws[col][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col, opts in (("H", RELEVANCE), ("I", EVENT_TYPES), ("J", SENTIMENT)):
            dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f"{col}2:{col}{n}")
            for cell in ws[col][1:]:
                cell.fill = yellow
        ws.freeze_panes = "B2"
    wb.save(GOLD / "team_sheet.xlsx")


def rewrite_xlsx():
    """Chỉ ghi lại file gán (vd. khi đổi hướng dẫn) — giữ nguyên mẫu và gold_key (kể cả cột distil_*)."""
    key = pd.read_parquet(GOLD / "team_key.parquet")
    art = pd.read_parquet(CLEAN / "articles_clean.parquet",
                          columns=["article_id", "published_at", "title", "description", "content"])
    tick = llm_article_labels()[["article_id", "tickers"]]
    smp = key[["gold_id", "article_id", "overlap", "src"]].merge(art, on="article_id", how="left").merge(tick, on="article_id", how="left")
    write_xlsx(smp)
    print(f"Đã ghi lại {GOLD / 'team_sheet.xlsx'} ({len(smp)} bài, mẫu không đổi)")


if __name__ == "__main__":
    import sys
    rewrite_xlsx() if "--xlsx-only" in sys.argv else main()
